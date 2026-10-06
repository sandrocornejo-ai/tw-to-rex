"""
excel_rapido.py — Lectura y escritura rápida de Excel.

* Lectura: usa python-calamine (motor en Rust, ~10x más rápido que openpyxl).
  Si no está instalado, usa openpyxl en modo solo lectura.
* Columnas nuevas en el archivo de entrada: edita el XML de la hoja
  directamente (segundos en vez de ~30 s con openpyxl). Si la hoja tiene
  algo que este método no sabe mover (fórmulas, tablas, imágenes, links…),
  retorna None y el llamador usa openpyxl.
* Escritura de la salida: usa XlsxWriter si está instalado.

Instalar los motores rápidos:  pip install python-calamine xlsxwriter
"""

import io
import re
import zipfile

import openpyxl

try:
    from python_calamine import CalamineWorkbook
except ImportError:                       # pragma: no cover
    CalamineWorkbook = None

try:
    import xlsxwriter
except ImportError:                       # pragma: no cover
    xlsxwriter = None


def motores_rapidos():
    """Lista de motores rápidos que faltan instalar (vacía si están todos)."""
    faltan = []
    if CalamineWorkbook is None:
        faltan.append('python-calamine')
    if xlsxwriter is None:
        faltan.append('xlsxwriter')
    return faltan


# ─────────────────────────────────────────────────────────────────
#  Lectura
# ─────────────────────────────────────────────────────────────────

def _norm(v):
    """Deja los valores de calamine igual que los de openpyxl."""
    if v == '':
        return None
    if v.__class__ is float and v.is_integer():
        return int(v)
    return v


def _calamine_wb(fuente):
    if isinstance(fuente, (bytes, bytearray)):
        return CalamineWorkbook.from_filelike(io.BytesIO(fuente))
    if hasattr(fuente, 'read'):
        return CalamineWorkbook.from_filelike(fuente)
    return CalamineWorkbook.from_path(str(fuente))


def _filas_calamine(ws):
    r0, c0 = ws.start if ws.start else (0, 0)
    for _ in range(r0):
        yield ()
    pad = (None,) * c0
    for fila in ws.iter_rows():
        yield pad + tuple(map(_norm, fila))


def leer_hoja(fuente, elegir=None):
    """Retorna (indice_hoja, total_filas, iterador de filas).

    Cada fila es una tupla de valores (celda vacía = None), igual que
    openpyxl con values_only=True.
    elegir(primeras_filas) -> bool: opcional; elige la primera hoja que
    cumpla (recibe las 8 primeras filas). Si ninguna cumple, usa la primera.
    """
    if CalamineWorkbook is not None:
        wb = _calamine_wb(fuente)
        hojas = [wb.get_sheet_by_index(i) for i in range(len(wb.sheet_names))]
        idx = 0
        if elegir and len(hojas) > 1:
            for i, ws in enumerate(hojas):
                it = _filas_calamine(ws)
                cab = [f for _, f in zip(range(8), it)]
                if elegir(cab):
                    idx = i
                    break
        ws = hojas[idx]
        r0 = ws.start[0] if ws.start else 0
        return idx, (ws.height or 0) + r0, _filas_calamine(ws)

    # respaldo: openpyxl
    if isinstance(fuente, (bytes, bytearray)):
        fuente = io.BytesIO(fuente)
    wb = openpyxl.load_workbook(fuente, read_only=True, data_only=True)
    idx = 0
    if elegir and len(wb.worksheets) > 1:
        for i, ws in enumerate(wb.worksheets):
            if elegir(list(ws.iter_rows(min_row=1, max_row=8, values_only=True))):
                idx = i
                break
    ws = wb.worksheets[idx]

    def gen():
        try:
            yield from ws.iter_rows(values_only=True)
        finally:
            wb.close()
    return idx, ws.max_row or 0, gen()


# ─────────────────────────────────────────────────────────────────
#  Insertar columnas editando el XML
# ─────────────────────────────────────────────────────────────────

_NO_SOPORTADO = ('<f>', '<f ', '<tableParts', '<hyperlinks', '<drawing', '<legacyDrawing',
                 '<oleObjects', '<controls', '<picture', '<sheetProtection', '<extLst')
_RE_CELDA = re.compile(r'<c r="([A-Z]{1,3})(\d+)"')
_RE_REF = re.compile(r'(\$?)([A-Z]{1,3})(\$?)(\d+)')
_RE_ATTR_REF = re.compile(r'\b(ref|sqref|activeCell|topLeftCell)="([^"]*)"')
_RE_FILA = re.compile(r'(<row\b[^>]*?\br="(\d+)"[^>]*?)(/>|>(.*?)</row>)', re.S)
_RE_COL = re.compile(r'<col\b([^>]*)/>')


def _col_num(letras):
    n = 0
    for ch in letras:
        n = n * 26 + ord(ch) - 64
    return n


def _col_letras(n):
    s = ''
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _xml_esc(t):
    return (str(t).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            .replace('"', '&quot;'))


def _ruta_hoja(z, indice):
    """Ruta dentro del zip de la hoja número `indice` (orden del libro)."""
    wbx = z.read('xl/workbook.xml').decode('utf-8')
    if '<definedName' in wbx:
        return None                      # nombres definidos: no se mueven
    hojas = re.findall(r'<sheet\b[^>]*?r:id="([^"]+)"', wbx)
    if indice >= len(hojas):
        return None
    rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    for m in re.finditer(r'<Relationship\b([^>]*)/>', rels):
        a = m.group(1)
        if f'Id="{hojas[indice]}"' in a:
            destino = re.search(r'Target="([^"]+)"', a).group(1)
            destino = destino.lstrip('/')
            return destino if destino.startswith('xl/') else 'xl/' + destino
    return None


def insertar_columnas(fuente, indice_hoja, col_ref, fila_titulos, titulos, valores,
                      ancho_min=14, avance=None):
    """Inserta len(titulos) columnas a la derecha de `col_ref` (1-based).

    fila_titulos: fila (1-based) donde van los títulos; toman el estilo de
                  la celda de `col_ref` en esa fila.
    valores: {fila (1-based): (v1, v2, …)}; números con formato #,##0,
             textos como texto.
    Retorna los bytes del nuevo archivo, o None si la hoja tiene algo que
    este método no sabe mover (el llamador debe usar openpyxl).
    """
    datos = fuente if isinstance(fuente, (bytes, bytearray)) else open(fuente, 'rb').read()
    k = len(titulos)
    zin = zipfile.ZipFile(io.BytesIO(datos))
    ruta = _ruta_hoja(zin, indice_hoja)
    if not ruta or ruta not in zin.namelist():
        return None
    xml = zin.read(ruta).decode('utf-8')
    if any(t in xml for t in _NO_SOPORTADO) or '<formula' in xml:
        return None
    if xml.count('<c ') != xml.count('<c r="'):   # celdas sin referencia al inicio
        return None
    if any(n.startswith(('xl/pivot', 'xl/tables', 'xl/charts', 'xl/drawings'))
           for n in zin.namelist()):
        return None
    if avance:
        avance(0.15)

    # ── estilo numérico #,##0 (numFmtId 3) en styles.xml ──
    estilos = zin.read('xl/styles.xml').decode('utf-8')
    m = re.search(r'<cellXfs count="(\d+)">', estilos)
    if not m:
        return None
    s_num = int(m.group(1))
    estilos = (estilos[:m.start()] + f'<cellXfs count="{s_num + 1}">' +
               estilos[m.end():].replace(
                   '</cellXfs>',
                   '<xf numFmtId="3" fontId="0" fillId="0" borderId="0" xfId="0" '
                   'applyNumberFormat="1"/></cellXfs>', 1))

    # ── mover referencias ──
    def mover_ref(mr):
        c = _col_num(mr.group(2))
        if c > col_ref:
            c += k
        return f"{mr.group(1)}{_col_letras(c)}{mr.group(3)}{mr.group(4)}"

    def mover_attr(ma):
        return f'{ma.group(1)}="{_RE_REF.sub(mover_ref, ma.group(2))}"'

    i_data = xml.index('<sheetData')
    f_data = xml.index('</sheetData>') if '</sheetData>' in xml else xml.index('/>', i_data) + 2
    cabeza, cuerpo, cola = xml[:i_data], xml[i_data:f_data], xml[f_data:]

    # cabeza: dimension, selección, panel y anchos de columna
    cabeza = _RE_ATTR_REF.sub(mover_attr, cabeza)
    ancho_ref = None
    cols = []
    for mc in _RE_COL.finditer(cabeza):
        a = mc.group(1)
        mn = int(re.search(r'\bmin="(\d+)"', a).group(1))
        mx = int(re.search(r'\bmax="(\d+)"', a).group(1))
        if mn <= col_ref <= mx:
            w = re.search(r'\bwidth="([\d.]+)"', a)
            ancho_ref = float(w.group(1)) if w else None
        tramos = []
        if mx <= col_ref:
            tramos.append((mn, mx))
        elif mn > col_ref:
            tramos.append((mn + k, mx + k))
        else:
            tramos += [(mn, col_ref), (col_ref + k + 1, mx + k)]
        for a_, b_ in tramos:
            b_ = min(b_, 16384)
            if a_ <= b_:
                cols.append((a_, re.sub(r'\bmin="\d+"', f'min="{a_}"',
                                        re.sub(r'\bmax="\d+"', f'max="{b_}"', a))))
    if '<cols>' in cabeza:
        ancho = max(ancho_ref or 0, ancho_min)
        for j in range(k):
            c = col_ref + 1 + j
            cols.append((c, f' min="{c}" max="{c}" width="{ancho}" customWidth="1"'))
        cols.sort()
        nuevo = '<cols>' + ''.join(f'<col{a}/>' for _, a in cols) + '</cols>'
        cabeza = re.sub(r'<cols>.*?</cols>', lambda _: nuevo, cabeza, count=1, flags=re.S)

    # cola: celdas combinadas, formato condicional, validaciones, filtros
    cola = _RE_ATTR_REF.sub(mover_attr, cola)

    # cuerpo: celdas de cada fila + columnas nuevas
    letras_nuevas = [_col_letras(col_ref + 1 + j) for j in range(k)]
    letra_ref = _col_letras(col_ref)
    total = max(len(cuerpo), 1)

    def celdas_nuevas(fila, vals, estilo_titulo=None):
        out = []
        for letra, v in zip(letras_nuevas, vals):
            r = f'{letra}{fila}'
            if estilo_titulo is not None or isinstance(v, str):
                s = f' s="{estilo_titulo}"' if estilo_titulo is not None else ''
                out.append(f'<c r="{r}"{s} t="inlineStr"><is><t>{_xml_esc(v)}</t></is></c>')
            elif v is not None:
                out.append(f'<c r="{r}" s="{s_num}"><v>{v}</v></c>')
        return ''.join(out)

    def mover_fila(mf):
        if avance and mf.start() % 2_000_000 < 400:
            avance(0.15 + 0.75 * mf.start() / total)
        inicio, fila = mf.group(1), int(mf.group(2))
        inicio = re.sub(r'\sspans="[^"]*"', '', inicio)
        contenido = mf.group(4) or ''
        contenido = _RE_CELDA.sub(
            lambda mc: f'<c r="{_col_letras(_col_num(mc.group(1)) + (k if _col_num(mc.group(1)) > col_ref else 0))}{mc.group(2)}"',
            contenido)
        if fila == fila_titulos:
            ms = re.search(rf'<c r="{letra_ref}{fila}"[^>]*?\ss="(\d+)"', contenido)
            extra = celdas_nuevas(fila, titulos, ms.group(1) if ms else '0')
        elif fila in valores:
            extra = celdas_nuevas(fila, valores[fila])
        else:
            extra = ''
        if extra:
            pos = len(contenido)
            for mc in _RE_CELDA.finditer(contenido):
                if _col_num(mc.group(1)) > col_ref:
                    pos = mc.start()
                    break
            contenido = contenido[:pos] + extra + contenido[pos:]
        if not contenido:
            return inicio + '/>'
        return f'{inicio}>{contenido}</row>'

    cuerpo = _RE_FILA.sub(mover_fila, cuerpo)
    filas_vistas = {int(x) for x in re.findall(r'<row\b[^>]*?\br="(\d+)"', cuerpo)}
    if fila_titulos not in filas_vistas or not set(valores) <= filas_vistas:
        return None                      # filas que no existen en el XML
    xml = cabeza + cuerpo + cola
    if avance:
        avance(0.92)

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED, compresslevel=1) as zout:
        for item in zin.infolist():
            if item.filename == ruta:
                zout.writestr(item.filename, xml)
            elif item.filename == 'xl/styles.xml':
                zout.writestr(item.filename, estilos)
            elif item.filename == 'xl/calcChain.xml':
                continue
            else:
                zout.writestr(item, zin.read(item.filename))
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────
#  Escritura de la salida
# ─────────────────────────────────────────────────────────────────

def escribir_filas(filas, nombre_hoja, formatos=None, avance=None):
    """Excel con `filas` (la primera es el encabezado) usando XlsxWriter.
    formatos: {indice_columna: función(fila) -> formato numérico o None}.
    Retorna bytes, o None si XlsxWriter no está instalado."""
    if xlsxwriter is None:
        return None
    formatos = formatos or {}
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {'constant_memory': True,
                                   'strings_to_numbers': False,
                                   'strings_to_formulas': False,
                                   'strings_to_urls': False})
    ws = wb.add_worksheet(nombre_hoja)
    cache = {}
    w_str, w_num = ws.write_string, ws.write_number
    total = max(len(filas), 1)
    for i, fila in enumerate(filas):
        if avance and i % 5000 == 0:
            avance(0.9 * i / total)
        for j, v in enumerate(fila):
            if v is None:
                continue
            fmt = None
            if i and j in formatos:
                nf = formatos[j](fila)
                if nf:
                    fmt = cache.get(nf) or cache.setdefault(nf, wb.add_format({'num_format': nf}))
            if v.__class__ is str:
                w_str(i, j, v, fmt)
            elif isinstance(v, bool):
                ws.write_boolean(i, j, v, fmt)
            elif isinstance(v, (int, float)):
                w_num(i, j, v, fmt)
            else:
                ws.write(i, j, v, fmt)
    wb.close()
    return buf.getvalue()
